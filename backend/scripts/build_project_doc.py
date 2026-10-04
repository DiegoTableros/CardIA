"""Genera docs/CardIA_Documentacion_Proyecto.docx (borrador completo del proyecto, sin formato elaborado).

Uso: uv run python -m scripts.build_project_doc
"""

import docx
from docx.enum.text import WD_BREAK

from app.core.config import REPO_DIR

OUT = REPO_DIR / "docs" / "CardIA_Documentacion_Proyecto.docx"
d = docx.Document()


def h(text: str, level: int = 1) -> None:
    d.add_heading(text, level=level)


def p(text: str) -> None:
    d.add_paragraph(text)


def bullets(items: list[str]) -> None:
    for i in items:
        d.add_paragraph(i, style="List Bullet")


def table(header: list[str], rows: list[list[str]]) -> None:
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, name in enumerate(header):
        t.rows[0].cells[i].text = name
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v
    d.add_paragraph()


# ------------------------------------------------------------------ portada
d.add_heading("CardIA", 0)
p("Plataforma de educación financiera sobre tarjetas de crédito en México")
p("Documentación técnica del proyecto")
p("Autor: [COMPLETAR]")
p("Programa / diplomado: [COMPLETAR]")
p("Profesor: [COMPLETAR]")
p("Fecha de entrega: [COMPLETAR]")
p("Repositorio: https://github.com/DiegoTableros/CardIA")
p("Aplicación en línea: https://cardia-wine.vercel.app")
d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ------------------------------------------------------------------ 1
h("1. Resumen")
p(
    "CardIA es una aplicación web de educación financiera que ayuda a una persona a conocer las tarjetas de crédito "
    "del mercado mexicano, entender sus costos y encontrar la que mejor se adapta a su perfil. Integra un modelo de "
    "aprendizaje no supervisado (clustering de tarjetas) con un asistente conversacional basado en un modelo de lenguaje "
    "(LLM) que consulta una base de datos mediante herramientas (tools) y agentes."
)
p(
    "La aplicación contesta tres preguntas: (1) qué tarjetas existen, con sus características, requisitos, beneficios y "
    "comisiones; (2) cómo funciona una tarjeta de crédito (pago mínimo, pago para no generar intereses, CAT, etc.); y "
    "(3) cuál es la mejor tarjeta para el usuario, a partir de un cuestionario corto."
)
p(
    "CardIA es una herramienta educativa. No está afiliada a ninguna institución financiera, no garantiza la aprobación "
    "de ningún crédito, no sugiere contratar productos y nunca solicita datos personales sensibles."
)

h("2. Objetivos y alcance")
h("2.1 Objetivo general", 2)
p(
    "Construir una aplicación desplegable que integre un modelo de machine learning con agentes de IA para orientar a usuarios sobre tarjetas de crédito en México."
)
h("2.2 Objetivos específicos", 2)
bullets(
    [
        "Consolidar una base de 69 tarjetas con costos, requisitos, comisiones y beneficios a partir de fuentes públicas (Banxico y CONDUSEF).",
        "Agrupar las tarjetas en perfiles mediante clustering y usar ese resultado para recomendar tarjetas y asignar un perfil al usuario.",
        "Ofrecer un asistente conversacional que use herramientas para consultar la base y nunca invente datos.",
        "Desplegar la aplicación en Vercel con una interfaz clara, moderna y accesible.",
    ]
)
h("2.3 Fuera de alcance (por diseño)", 2)
bullets(
    [
        "Scraping, conexión con instituciones financieras o ejecución de operaciones.",
        "Datos personales sensibles (RFC, CURP, cuentas, números de tarjeta).",
        "Cualquier sugerencia de contratar un producto. El enlace «Aplica ahora» lleva al sitio general de la institución, nunca a una URL de producto.",
    ]
)

# ------------------------------------------------------------------ 3
h("3. Arquitectura")
p(
    "El proyecto es un monorepo con dos aplicaciones: un frontend Angular y un backend FastAPI. El backend consulta una "
    "base SQLite que se siembra a partir de archivos JSON generados desde el Excel de tarjetas, ejecuta el recomendador "
    "basado en los perfiles (clusters) y orquesta el asistente con el SDK oficial de OpenAI."
)
table(
    ["Capa", "Tecnología", "Responsabilidad"],
    [
        [
            "Frontend",
            "Angular 22 (standalone + signals), Tailwind CSS 4",
            "Interfaz de usuario y administración",
        ],
        ["Backend", "Python 3.12, FastAPI, Pydantic v2", "API REST, reglas de negocio, agentes"],
        [
            "Datos",
            "SQLite (SQLAlchemy 2 async + aiosqlite)",
            "Tarjetas, comisiones, beneficios, usuarios, trazas",
        ],
        ["ML", "Clustering jerárquico (Ward) + recomendador propio", "Perfiles de tarjetas y ranking"],
        ["LLM", "SDK oficial de OpenAI (Responses API)", "Planificación y redacción de respuestas"],
        [
            "Contratos",
            "OpenAPI → TypeScript (openapi-typescript)",
            "Tipos del frontend generados desde la API",
        ],
        ["Despliegue", "Vercel (Services: frontend + backend)", "Un solo proyecto y un solo dominio"],
    ],
)
p(
    "Flujo general: el usuario interactúa con el frontend; este llama a la API con un token JWT; la API consulta la base, "
    "ejecuta el recomendador o el asistente y devuelve respuestas tipadas con Pydantic. El asistente usa el LLM únicamente "
    "para decidir qué herramientas consultar y para redactar la respuesta con los datos que esas herramientas devuelven."
)

# ------------------------------------------------------------------ 4
h("4. Datos")
h("4.1 Fuentes", 2)
table(
    ["Fuente", "Contenido"],
    [
        [
            "tarjetas.xlsx",
            "69 tarjetas. Hojas: Tarjetas, Requisitos, Comisiones, Beneficios, Instituciones y dataset_tarjetas_ingenieria. Todo se relaciona por ID_Tarjeta. La columna Cluster (0 a 5) es el resultado del modelo.",
        ],
        ["imagenes_tarjetas/", "Una imagen por tarjeta (el nombre del archivo es el nombre de la tarjeta)."],
        [
            "Glosario CONDUSEF",
            "58 términos (generales, seguros y beneficios), transcritos del PDF oficial porque este es un documento escaneado.",
        ],
        [
            "Comisiones ERCO (Banxico)",
            "Archivo descargado manualmente; está identificado para una integración futura (se cruza por institución y nombre del producto).",
        ],
    ],
)
h("4.2 Proceso de datos (ETL)", 2)
bullets(
    [
        "Los encabezados del Excel se normalizan a snake_case sin acentos en una sola función compartida.",
        "make data convierte el Excel en backend/data/cards.json (tarjetas, comisiones, beneficios, URL de la institución, imagen, cluster y las variables de ingeniería del modelo).",
        "make images procesa las imágenes: quita el fondo, recorta bordes y las coloca en un lienzo horizontal transparente de 560×353 px (las verticales quedan más pequeñas, no más largas) y las guarda como webp por ID.",
        "El backend nunca lee el Excel en tiempo de petición: solo lee los artefactos generados.",
        "Al arrancar, la API siembra SQLite de forma idempotente. Si cambia la versión de los datos (campo generated_at), se recrean solo las tablas estáticas; usuarios y trazas no se tocan.",
    ]
)
h("4.3 Valores faltantes", 2)
p(
    "Hay requisitos sin dato publicado (por ejemplo, score faltante en 31 de 69 tarjetas). No se imputan en silencio: la interfaz los muestra como «No publicado» y el recomendador marca la tarjeta como «por confirmar»."
)

# ------------------------------------------------------------------ 5
h("5. Perfiles de tarjetas y recomendador")
p(
    "El detalle del modelo de clustering (variables, método y análisis de cada grupo) está en el documento de "
    "perfilamiento de clusters. Aquí se describe cómo se integra en la aplicación."
)
h("5.1 Los seis perfiles", 2)
table(
    ["Cluster", "Nombre", "Característica principal"],
    [
        [
            "0",
            "Historial sólido",
            "Requisitos exigentes (score cercano a 700, 12 meses de antigüedad) con descuentos amplios",
        ],
        ["1", "Costo bajo", "Anualidad más baja e ingreso mínimo accesible"],
        ["2", "Viajes y recompensas", "Puntos por dólar, MSI de viaje y preventas"],
        ["3", "Promociones", "Mayor variedad de beneficios, sin pedir score; CAT más alto"],
        ["4", "Comisiones en pesos", "Pocos beneficios y comisiones fijas altas"],
        ["5", "Tasa baja", "Menor tasa y CAT, con anualidad alta"],
    ],
)
p("Los nombres y textos viven en backend/data/clusters.json y se pueden editar sin modificar código.")
h("5.2 Cuestionario", 2)
p(
    "El usuario responde sobre: edad e ingreso mensual (en rangos), historial (score, antigüedad laboral y residencial), uso principal de la tarjeta, forma de pago (totalero, a veces financia, suele financiar), si desea evitar anualidad, si desea evitar comisiones altas y los beneficios de interés (puede elegir varios). No se solicita ningún dato personal."
)
h("5.3 Método de recomendación", 2)
bullets(
    [
        "Elegibilidad: se descartan las tarjetas cuyos requisitos publicados no se cumplen; si falta algún dato queda «por confirmar» (puntaje −8 %).",
        "Utilidad: cada tarjeta se mide en 9 dimensiones entre 0 y 1 (anualidad, tasa, comisiones, puntos, viaje, promociones, MSI, transferencia de saldo y requisitos accesibles), escaladas por percentiles 5–95. Las respuestas del usuario definen el peso de cada dimensión.",
        "Puntaje de tarjeta = 75 % utilidad propia + 25 % afinidad de su cluster (utilidad media de las tarjetas del cluster ajustada por elegibilidad). Si el usuario evita anualidad, las tarjetas con anualidad se multiplican por 0.75.",
        "Orden: Top 1, Top 2, etc. Al usuario no se le muestra el puntaje porque no le es interpretable.",
        "Perfil del usuario = cluster de su tarjeta Top 1.",
        "Advertencias (informan, no descartan): la tarjeta contradice lo que pidió (por ejemplo, cobra anualidad), tasa o CAT altos si no paga el total, comisiones altas si las quiere evitar, o falta un beneficio buscado. Las tarjetas atípicas de su cluster solo se advierten.",
    ]
)
h("5.4 Vista Perfiles (administrador)", 2)
p(
    "Muestra los seis perfiles con sus promedios, un diagrama del proceso de asignación, un mapa de anualidad contra tasa, distribuciones del mercado y la documentación completa de cada cluster tomada del documento de perfilamiento."
)

# ------------------------------------------------------------------ 6
h("6. Backend")
h("6.1 Estructura", 2)
table(
    ["Carpeta", "Contenido"],
    [
        ["app/api/v1", "Routers: auth, cards, compare, recommend, chat, bi, admin, meta"],
        ["app/agents", "Planner, orquestador, narrador y capa LLM"],
        ["app/tools", "Registro de 8 tools que consultan la base"],
        ["app/domain", "Modelos de dominio, normalización, base educativa y glosario"],
        ["app/ml", "Perfiles (clusters) y recomendador"],
        ["app/services", "Catálogo en memoria, comparador, BI, auditoría"],
        ["app/schemas", "Modelos Pydantic v2 de entrada y salida"],
        ["app/db", "Modelos SQLAlchemy, sesión y seed"],
        ["app/core", "Configuración, seguridad (JWT y argon2), privacidad, errores y aviso"],
        ["scripts", "ETL de datos, procesado de imágenes, reporte de clusters y exportación de OpenAPI"],
        ["tests", "Pruebas automatizadas"],
    ],
)
h("6.2 API", 2)
p(
    "Prefijo /api/v1. Todo requiere token Bearer excepto meta/* y auth/login. Los endpoints de BI/Perfiles y administración exigen rol administrador."
)
table(
    ["Método", "Ruta", "Descripción"],
    [
        ["GET", "/meta/health, /meta/disclaimer", "Estado del servicio y aviso"],
        ["POST/GET", "/auth/login, /auth/me, /auth/logout", "Sesión con JWT"],
        ["GET", "/cards, /cards/facets, /cards/{id}", "Catálogo, filtros y detalle"],
        ["POST", "/compare", "Compara de 2 a 4 tarjetas"],
        ["POST", "/recommend", "Perfil y Top de tarjetas"],
        ["GET", "/chat/topics", "Temas para aprender (glosario)"],
        ["POST", "/chat/plan, /chat/runs/{id}/execute, /chat/ask", "Asistente"],
        ["GET", "/chat/runs, /chat/runs/{id}", "Historial de conversaciones"],
        ["GET", "/bi", "Perfiles y documentación (admin)"],
        ["GET", "/admin/overview, events, chat-runs, users", "Trazabilidad (admin)"],
    ],
)
h("6.3 Seguridad y trazabilidad", 2)
bullets(
    [
        "Contraseñas con argon2 y sesiones con JWT firmado.",
        "Roles: usuario y administrador. Perfiles y Administración son solo para administrador, tanto en la interfaz como en la API.",
        "Cada inicio de sesión, recomendación, comparación y conversación del asistente se registra como evento de auditoría; el administrador los consulta en el panel.",
        "El chat enmascara RFC, CURP, números de tarjeta, correos y teléfonos antes de guardarlos o enviarlos al LLM.",
        "En producción el backend no arranca si JWT_SECRET falta o tiene menos de 32 caracteres.",
    ]
)

# ------------------------------------------------------------------ 7
h("7. Asistente con IA (agentes y herramientas)")
h("7.1 Principio de diseño", 2)
p(
    "El LLM nunca calcula ni inventa datos de tarjetas. Toda cifra (anualidad, tasa, requisitos, beneficios, comisiones) proviene de una tool que consulta la base; el cluster y las recomendaciones provienen del recomendador, no del LLM."
)
h("7.2 Flujo", 2)
bullets(
    [
        "El orquestador produce un plan explícito de tool calls con dependencias antes de ejecutar y lo guarda en la base para trazabilidad.",
        "Las tools se ejecutan por niveles de dependencia.",
        "El narrador redacta la respuesta en markdown usando únicamente los resultados de las tools.",
        "Al usuario final solo se le muestran puntitos de color de los agentes que responden; el plan no se muestra.",
    ]
)
h("7.3 Agentes y tools", 2)
table(
    ["Agente", "Responsabilidad", "Tools"],
    [
        [
            "FundamentalsAgent",
            "Información de tarjetas de la base",
            "search_cards, get_card_details, compare_cards",
        ],
        ["ComisionesAgent", "Desglose de comisiones e impacto si se tiene contratada", "get_card_fees"],
        ["EducativeAgent", "Explicaciones de inclusión financiera", "get_education_topic, search_glossary"],
        ["PerfilAgent", "Interpreta los perfiles (clusters)", "explain_profiles, get_card_profile"],
    ],
)
h("7.4 Integración con OpenAI", 2)
bullets(
    [
        "SDK oficial de OpenAI con la Responses API.",
        "Planner: salida estructurada (Pydantic). El plan se valida contra el registro de tools y el catálogo antes de aceptarse.",
        "Narrador: modelo gpt-5.4-mini, verbosidad media y razonamiento bajo (configurables por variables de entorno).",
        "Respaldo: si no hay clave o OpenAI falla, el planner y el narrador por reglas responden con las mismas tools. El usuario no ve el error.",
        "Se usa el historial reciente de la sesión para resolver referencias como «esa tarjeta».",
        "Hallazgo: el modelo gpt-5.1-mini configurado al inicio no existe en la cuenta; por eso el asistente caía siempre en el modo de reglas. Se verificó la lista de modelos disponibles y se cambió a gpt-5.4-mini.",
    ]
)
h("7.5 Educación financiera", 2)
p(
    "El agente educativo se apoya en el glosario oficial de CONDUSEF (58 términos) y en temas redactados a mano (qué es una tarjeta, fecha de corte, pago mínimo, pago para no generar intereses, CAT, tasa, anualidad, MSI, disposición de efectivo, buró y score, comisiones, beneficios, transferencia de saldo y riesgos). Los ejemplos numéricos son ilustrativos y se indican como tales."
)

# ------------------------------------------------------------------ 8
h("8. Frontend")
h("8.1 Tecnologías y estructura", 2)
bullets(
    [
        "Angular 22 con componentes standalone y signals (sin NgModules).",
        "Tailwind CSS 4; el sistema de diseño (colores, tipografías, componentes base) vive en un solo archivo de estilos con tokens.",
        "Estructura src/app: core (API, autenticación, formato), shared (componentes reutilizables), features (una carpeta por pantalla) y viz (gráficos).",
        "Gráficos propios en SVG (barras, dona y dispersión), sin librerías de gráficas.",
        "Los tipos de la API se generan desde /openapi.json con openapi-typescript (make contracts).",
        "Carga diferida (lazy loading) de cada vista; el bundle inicial pesa unos 417 kB (103 kB comprimido).",
    ]
)
h("8.2 Vistas", 2)
table(
    ["Vista", "Ruta", "Acceso", "Descripción"],
    [
        [
            "Inicio de sesión",
            "/login",
            "Público",
            "Acceso con usuario y contraseña; muestra tarjetas reales y el mensaje de la herramienta.",
        ],
        [
            "Inicio",
            "/",
            "Usuario",
            "Presentación, tarjetas destacadas (tasa baja, sin anualidad, menos comisiones), invitación al asistente y perfiles.",
        ],
        [
            "Encuentra tu tarjeta",
            "/encuentra-tu-tarjeta",
            "Usuario",
            "Cuestionario por pasos; resultado con perfil, Top de tarjetas, advertencias y «Tus respuestas».",
        ],
        [
            "Catálogo",
            "/tarjetas",
            "Usuario",
            "Filtros fijos a la izquierda y tarjetas con imagen de dos en dos.",
        ],
        [
            "Detalle de tarjeta",
            "/tarjetas/:id",
            "Usuario",
            "Requisitos, beneficios, comisiones en tres columnas (obligatorias, por evento, penalizaciones) y «Aplica ahora».",
        ],
        [
            "Comparar",
            "/comparar",
            "Usuario",
            "Comparador de 2 a 4 tarjetas que resalta el mejor valor de cada métrica.",
        ],
        ["Asistente", "/asistente", "Usuario", "Chat con IA, temas para aprender y tipos de pregunta."],
        [
            "Perfiles",
            "/perfiles",
            "Administrador",
            "Perfiles, método de asignación, mapa y documentación de los clusters.",
        ],
        [
            "Administración",
            "/admin",
            "Administrador",
            "KPIs, eventos de trazabilidad, ejecuciones del asistente y usuarios.",
        ],
    ],
)
h("8.3 Experiencia de usuario", 2)
bullets(
    [
        "Producto de cara al usuario final: no se muestran conteos de la base ni menciones a machine learning o agentes (salvo en Perfiles, para administradores).",
        "Un solo aviso de educación financiera, igual en todas las páginas, y pie de página con las fuentes.",
        "Estados de carga, error y vacío en cada vista; diseño responsive (mobile first); etiquetas en formularios y foco visible.",
        "Las imágenes reales de las tarjetas se muestran en un escenario uniforme; si una imagen falla se usa una ilustración generada.",
        "Los guards de rutas devuelven UrlTree para no dejar pantallas en blanco.",
    ]
)

# ------------------------------------------------------------------ 9
h("9. Calidad y pruebas")
bullets(
    [
        "31 pruebas de backend con pytest, sin llamadas de red (OpenAI se sustituye por un cliente falso): normalización, ETL, recomendador, tools, planner, API, privacidad y capa LLM.",
        "3 pruebas de frontend con vitest (renderizador de markdown seguro y formato).",
        "Una prueba reproduce las cifras del documento de perfilamiento (tamaño, CAT y anualidad promedio por cluster) desde los datos cargados.",
        "Lint con ruff (backend) y verificación de tipos con tsc (frontend).",
        "Pruebas de humo manuales contra la API y verificación del arranque en modo producción simulado.",
    ]
)

# ------------------------------------------------------------------ 10
h("10. Despliegue en Vercel")
h("10.1 Estrategia", 2)
p(
    "Se usa la función Services de Vercel para desplegar frontend y backend como un solo proyecto en un solo dominio. El archivo vercel.json define dos servicios y dos reglas de ruteo."
)
p(
    "Servicio frontend: carpeta frontend, framework Angular (sitio estático). Servicio backend: carpeta backend, framework FastAPI, punto de entrada app.main:app."
)
p(
    "Reglas de ruteo (se evalúan en orden): /api(/.*)? va al backend; cualquier otra ruta va al frontend. El backend recibe la ruta original (/api/v1/...), igual que en local, y el frontend usa rutas relativas, por lo que no requiere CORS ni variables de entorno."
)
h("10.2 Variables de entorno en Vercel", 2)
table(
    ["Variable", "Obligatoria", "Descripción"],
    [
        ["OPENAI_API_KEY", "Para activar el LLM", "Clave de OpenAI (sin ella el asistente usa reglas)"],
        [
            "JWT_SECRET",
            "Sí",
            "Secreto de firma de sesiones, 32 o más caracteres; sin él el backend no arranca",
        ],
        ["SEED_ADMIN_PASSWORD", "Recomendada", "Contraseña de la cuenta administrador"],
        ["SEED_DEMO_PASSWORD", "Recomendada", "Contraseña de la cuenta de demostración"],
        ["LLM_MODEL, LLM_MODEL_FAST", "No", "Por defecto gpt-5.4-mini"],
    ],
)
h("10.3 Pasos de despliegue", 2)
bullets(
    [
        "Crear el proyecto en Vercel e importar el repositorio de GitHub (proyecto CardIA).",
        "Definir las variables de entorno de la sección 10.2.",
        "Confirmar que el proyecto usa Services (vercel.json en la raíz) y la versión de Node 24.x.",
        "Hacer push a la rama main; Vercel construye y despliega automáticamente.",
        "Verificar https://cardia-wine.vercel.app/api/v1/meta/health (debe responder status ok y 69 tarjetas) y entrar con las cuentas semilla.",
        "[COMPLETAR] Registrar aquí la fecha, el resultado y cualquier ajuste realizado durante el despliegue.",
    ]
)
h("10.4 Limitaciones del despliegue", 2)
bullets(
    [
        "SQLite es efímero en Vercel (decisión D1): vive en /tmp y se re-siembra en cada arranque en frío. Los usuarios creados, las trazas y los planes del chat se pierden al reiniciar. Es aceptable para una demostración.",
        "Dos peticiones consecutivas pueden caer en instancias distintas: por eso existe /chat/ask (plan y ejecución en una sola petición) como respaldo automático del frontend.",
        "El historial del chat se guarda también en el navegador (localStorage).",
        "Las cuentas semilla deben protegerse con contraseñas propias definidas por variables de entorno.",
    ]
)

# ------------------------------------------------------------------ 11
h("11. Decisiones de arquitectura")
table(
    ["ID", "Decisión"],
    [
        ["D1", "SQLite efímero en Vercel; en local y Docker, SQLite en archivo."],
        ["D2", "Backend ligero: dependencias mínimas en runtime por el límite de tamaño de la función."],
        ["D3", "El LLM nunca calcula ni inventa datos de tarjetas; todo viene de tools o del recomendador."],
        [
            "D4",
            "Pipeline de recomendación reproducible y llamable: recommend(perfil) devuelve perfil, tarjetas, advertencias y supuestos.",
        ],
        ["D5", "Contrato de API estable con Pydantic v2; el frontend consume tipos generados."],
        ["D6", "El aviso de educación financiera aparece siempre."],
        ["D7", "Sin secretos en el repositorio; las claves viven en .env y en Vercel."],
        [
            "D8",
            "Plan explícito antes de ejecutar; se guarda para trazabilidad pero no se muestra al usuario final.",
        ],
        ["D9", "Un solo modelo: el perfil del usuario es el cluster de su tarjeta Top 1."],
        ["D10", "Sin cashback en el formulario (el Excel no tiene ese beneficio)."],
        ["D11", "Alcance prohibido: scraping, conexión con bancos, operaciones, datos sensibles."],
        ["D12", "Producto, no proyecto académico: sin conteos ni menciones de ML para el usuario final."],
        ["D13", "Aviso único con el mismo texto en frontend y backend."],
    ],
)

# ------------------------------------------------------------------ 12
h("12. Lecciones técnicas y problemas resueltos")
bullets(
    [
        "openapi-typescript declara dependencia de TypeScript 5 y choca con TypeScript 6 de Angular 22: se ejecuta con npx desde make contracts en lugar de instalarlo como dependencia.",
        "Tailwind 4: una clase reutilizada con @apply debe declararse con @utility y no dentro de @layer components.",
        "La recomendación inicial por reglas dejaba fuera a las tarjetas de $0 cuando el usuario no quería anualidad; se corrigió con mayor peso y penalización de la anualidad y se añadió una prueba.",
        "El planner por reglas clasificaba mal las preguntas «¿Qué es X?»; se normalizó el texto (signos de apertura) y se dio prioridad al glosario.",
        "El PDF del glosario es un escaneo: se transcribió manualmente a JSON.",
        "El procesado de imágenes con relleno de Pillow era demasiado lento; se reescribió con numpy.",
        "uvicorn con --reload deja un proceso hijo con la configuración anterior: para recargar el .env hay que cerrar también ese proceso.",
        "El modelo gpt-5.1-mini no existía en la cuenta; se verificó la lista de modelos y se usa gpt-5.4-mini.",
    ]
)

# ------------------------------------------------------------------ 13
h("13. Limitaciones y trabajo futuro")
bullets(
    [
        "Integrar las comisiones del portal ERCO de Banxico para un desglose más detallado (el archivo ya está identificado).",
        "Reentrenar el modelo de clustering cuando cambie la base y documentar el notebook en la carpeta notebooks.",
        "Afinar el recomendador (pesos y reglas de advertencia) con retroalimentación de usuarios.",
        "Persistencia durable (base de datos externa) si se requiere conservar usuarios y trazas en producción.",
        "Contenedores (Docker y docker compose) para ejecución local reproducible.",
    ]
)

# ------------------------------------------------------------------ 14
h("14. Cómo reproducir el proyecto")
p("Requisitos: Python 3.12, uv, Node 24 LTS y make.")
bullets(
    [
        "make setup — instala dependencias de backend y frontend.",
        "make dev — levanta la API en el puerto 8000 y la interfaz en el 4200.",
        "make test — ejecuta las pruebas de backend y frontend.",
        "make lint — verifica estilo y tipos.",
        "make contracts — regenera los tipos de TypeScript desde la API.",
        "make images y make data — regeneran las imágenes y los datos desde las fuentes.",
        "Para activar el asistente con IA, crear el archivo .env con OPENAI_API_KEY y reiniciar el backend.",
    ]
)

h("Anexo A. Cuentas de demostración (entorno local)")
p(
    "Administrador: admin@cardia.local. Usuario: demo@cardia.local. Las contraseñas locales por defecto están en el archivo .env.example del repositorio; en producción se definen con variables de entorno."
)
h("Anexo B. Referencias")
bullets(
    [
        "Banxico: información de tarjetas de crédito y portal ERCO.",
        "CONDUSEF: sitio de tarjetas de crédito y glosario oficial.",
        "Documentación oficial de Angular, FastAPI, SQLAlchemy, Pydantic, OpenAI y Vercel.",
        "Documento «perfilamiento_clusters» (descripción completa de los seis clusters).",
        "[COMPLETAR] Referencias adicionales del diplomado.",
    ]
)

OUT.parent.mkdir(parents=True, exist_ok=True)
d.save(str(OUT))
print(f"OK -> {OUT}")

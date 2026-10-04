"""ETL: datos_consolidados/tarjetas.xlsx -> backend/data/cards.json.

El backend NUNCA lee el Excel en tiempo de request; solo lee el JSON generado aqui.
Uso: uv run python -m scripts.build_data
"""

import json
import sys
from datetime import UTC, datetime

from openpyxl import load_workbook

from app.core.config import DATA_DIR, REPO_DIR
from app.domain.normalize import (
    clean_benefit_text,
    denomination_key,
    fee_type_key,
    normalize_header,
    to_float,
    to_int,
)

SOURCE = REPO_DIR / "datos_consolidados" / "tarjetas.xlsx"
OUTPUT = DATA_DIR / "cards.json"
IMAGES = DATA_DIR / "images.json"  # generado por scripts.build_images


def read_sheet(wb, name: str) -> list[dict[str, object]]:
    rows = wb[name].iter_rows(values_only=True)
    headers = [normalize_header(h) for h in next(rows)]
    out = []
    for row in rows:
        if all(v is None for v in row):
            continue
        out.append(dict(zip(headers, row, strict=False)))
    return out


def card_id(value: object) -> str:
    return str(value).strip().zfill(3)


def build() -> dict[str, object]:
    lock = SOURCE.parent / f"~${SOURCE.name}"
    if lock.exists():
        sys.exit(f"El Excel esta abierto ({lock.name}). Cierralo y vuelve a ejecutar.")

    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    cards_raw = read_sheet(wb, "Tarjetas")
    institutions = {
        str(i["institucion"]).strip(): (str(i["url"]).strip() if i.get("url") else None)
        for i in read_sheet(wb, "Instituciones")
        if i.get("institucion")
    }
    images = json.loads(IMAGES.read_text(encoding="utf-8")) if IMAGES.exists() else {}
    # Variables de ingenieria del modelo de clustering (beneficios y comisiones), por ID_Tarjeta.
    eng = {
        card_id(r["id_tarjeta"]): {k: to_float(v) for k, v in r.items() if k.startswith(("benef_", "com_"))}
        for r in read_sheet(wb, "dataset_tarjetas_ingenieria")
    }
    reqs = {card_id(r["id_tarjeta"]): r for r in read_sheet(wb, "Requisitos")}
    fees: dict[str, list[dict[str, object]]] = {}
    for f in read_sheet(wb, "Comisiones"):
        fees.setdefault(card_id(f["id_tarjeta"]), []).append(
            {
                "concept": str(f["concepto"]).strip(),
                "amount": to_float(f["monto"]),
                "denomination": denomination_key(f["denominacion"]),
                "fee_type": fee_type_key(f["tipo"]),
            }
        )
    benefits: dict[str, list[dict[str, object]]] = {}
    for b in read_sheet(wb, "Beneficios"):
        btype = str(b["tipo"]).strip()
        benefits.setdefault(card_id(b["id_tarjeta"]), []).append(
            {"benefit_type": btype, "text": clean_benefit_text(btype, str(b["texto"] or ""))}
        )

    cards = []
    for c in cards_raw:
        cid = card_id(c["id_tarjeta"])
        r = reqs.get(cid, {})
        institution = str(c["institucion"]).strip()
        image = images.get(cid) or {}
        cards.append(
            {
                "id": cid,
                "name": str(c["nombre_de_la_tarjeta"]).strip(),
                "institution": institution,
                "institution_url": institutions.get(institution),
                "image_url": image.get("file"),
                "image_orientation": image.get("orientation"),
                "cluster": to_int(c.get("cluster")),
                "features": eng.get(cid, {}),
                "card_class": str(c["clase"]).strip(),
                "cat": to_float(c["cat_publicidad"]),
                "annual_fee": to_float(c["anualidad"]),
                "interest_rate": to_float(c["tasa_de_interes_de_contrato"]),
                "credit_line_min": to_float(c["linea_de_credito_desde"]),
                "requirements": {
                    "age_min": to_int(r.get("edad_minima")),
                    "age_max": to_int(r.get("edad_maxima")),
                    "score_min": to_int(r.get("score_de_credito_minimo")),
                    "work_seniority_min": to_int(r.get("antiguedad_laboral_minima")),
                    "residence_seniority_min": to_int(r.get("antiguedad_residencial_minima")),
                    "monthly_income_min": to_float(r.get("ingreso_mensual_minimo")),
                },
                "fees": fees.get(cid, []),
                "benefits": benefits.get(cid, []),
            }
        )
    cards.sort(key=lambda x: x["id"])
    return {
        "source": "datos_consolidados/tarjetas.xlsx",
        "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "count": len(cards),
        "cards": cards,
    }


def main() -> None:
    data = build()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    n_fees = sum(len(c["fees"]) for c in data["cards"])
    n_ben = sum(len(c["benefits"]) for c in data["cards"])
    n_img = sum(1 for c in data["cards"] if c["image_url"])
    no_cluster = [c["id"] for c in data["cards"] if c["cluster"] is None]
    if no_cluster:
        print(f"ATENCION: tarjetas sin Cluster: {no_cluster}")
    no_url = sorted({c["institution"] for c in data["cards"] if not c["institution_url"]})
    print(
        f"OK {data['count']} tarjetas, {n_fees} comisiones, {n_ben} beneficios, {n_img} imagenes -> {OUTPUT}"
    )
    if no_url:
        print(f"Instituciones sin URL: {no_url}")


if __name__ == "__main__":
    main()

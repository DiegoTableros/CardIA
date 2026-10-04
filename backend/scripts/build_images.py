"""Optimiza las imagenes de tarjetas: imagenes_tarjetas/<Nombre de la tarjeta>.* -> frontend/public/cards/<ID>.webp

Todas salen en el MISMO lienzo horizontal transparente (560x353, proporcion de tarjeta 1.586):
- Se quita el fondo (blanco o transparente) conectado a los bordes de la imagen.
- Se recorta a la tarjeta, se erosiona el borde ~1.5 px y se limpia el color de la franja del borde
  (sin halo blanco) antes de suavizar la transparencia.
- Las tarjetas verticales se escalan para caber en la altura del lienzo (quedan mas chicas, no mas largas).
- Se escribe backend/data/images.json con {id: {file, orientation}} que usa el ETL.
Uso: uv run python -m scripts.build_images
"""

import json
import sys
import unicodedata
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFilter

from app.core.config import DATA_DIR, REPO_DIR
from app.domain.normalize import normalize_header

SRC_DIR = REPO_DIR / "imagenes_tarjetas"
OUT_DIR = REPO_DIR / "frontend" / "public" / "cards"
MANIFEST = DATA_DIR / "images.json"
EXCEL = REPO_DIR / "datos_consolidados" / "tarjetas.xlsx"

CANVAS = (560, 353)  # 1.586 = proporcion ISO/IEC 7810 ID-1
PAD = 8
WORK_MAX = 700  # lado maximo de trabajo (acelera el relleno de fondo)
SS = 2  # supersampling para bordes suaves
ERODE = 3  # px a escala SS (~1.5 px finales)
WHITE_MIN = 238  # canal minimo para considerar "fondo blanco"
MAX_BG_FRACTION = 0.7  # si el "fondo" supera esto, la tarjeta es blanca: no se quita nada
ALPHA_SOLID = 200  # por debajo es transparente/sombra (no es la tarjeta)
FRAME_DEPTH = 0.02  # profundidad maxima del marco claro a pelar, como fraccion del lado corto
FRAME_MIN = 185  # canal minimo para considerar un pixel "claro"
FRAME_SAT = 45  # diferencia max-min maxima para considerarlo poco saturado (blanco/gris)


def _key(text: str) -> str:
    return unicodedata.normalize("NFC", text).strip().casefold()


def _flood_from_border(cand: np.ndarray) -> np.ndarray:
    """Region de `cand` conectada a los bordes de la imagen."""
    border = np.zeros_like(cand)
    border[0, :] = border[-1, :] = border[:, 0] = border[:, -1] = True
    return _grow(cand & border, cand, 10_000)


def _grow(seed: np.ndarray, allowed: np.ndarray, steps: int) -> np.ndarray:
    """Dilatacion 4-conexa de `seed` restringida a `allowed` (hasta `steps` pasos o converger)."""
    cur = seed.copy()
    for _ in range(steps):
        g = cur.copy()
        g[1:] |= cur[:-1]
        g[:-1] |= cur[1:]
        g[:, 1:] |= cur[:, :-1]
        g[:, :-1] |= cur[:, 1:]
        g &= allowed
        if (g == cur).all():
            break
        cur = g
    return cur


def _peel_light_frame(bg: np.ndarray, rgb: np.ndarray, depth: int) -> np.ndarray:
    """Pela desde afuera hacia adentro una franja clara y poco saturada pegada al fondo (marco blanquito)."""
    mx, mn = rgb.max(axis=2).astype(int), rgb.min(axis=2).astype(int)
    lightish = (mn >= FRAME_MIN) & (mx - mn <= FRAME_SAT)
    return _grow(bg, bg | lightish, depth)


def background_mask(rgba: np.ndarray) -> np.ndarray:
    """True donde hay fondo conectado a los bordes.

    - Fondo = transparente o semitransparente (sombras suaves) y blanco, conectado a los bordes.
    - Si "quitar el blanco" se come casi toda la tarjeta, es una tarjeta blanca sin contorno: solo se
      quita lo transparente y no se pela nada.
    - Despues se pela un marco claro de profundidad limitada.
    """
    alpha, rgb = rgba[..., 3], rgba[..., :3]
    see_through = alpha < ALPHA_SOLID
    transparent = _flood_from_border(see_through)
    with_white = _flood_from_border(see_through | (rgb.min(axis=2) >= WHITE_MIN))
    opaque = (~transparent).sum()
    if (with_white & ~transparent).sum() > MAX_BG_FRACTION * max(1, opaque):
        return transparent
    depth = max(2, round(FRAME_DEPTH * min(alpha.shape)))
    return _peel_light_frame(with_white, rgb, depth)


def _round_corners(mask: np.ndarray) -> np.ndarray:
    """Si la mascara es un rectangulo recto (esquinas llenas), redondea esquinas como una tarjeta ID-1."""
    h, w = mask.shape
    if not (mask[1, 1] and mask[1, w - 2] and mask[h - 2, 1] and mask[h - 2, w - 2]):
        return mask
    r = max(2, round(0.059 * min(w, h)))
    rounded = Image.new("L", (w, h), 0)
    ImageDraw.Draw(rounded).rounded_rectangle((0, 0, w - 1, h - 1), radius=r, fill=255)
    return mask & (np.asarray(rounded) > 127)


def _blur(arr: np.ndarray, sigma: float) -> np.ndarray:
    """Desenfoque gaussiano separable en numpy (Pillow no lo soporta para floats)."""
    r = max(1, int(3 * sigma))
    x = np.arange(-r, r + 1, dtype="float32")
    k = np.exp(-(x**2) / (2 * sigma**2))
    k /= k.sum()
    out = arr.astype("float32")
    for axis in (0, 1):
        pad = [(0, 0), (0, 0)]
        pad[axis] = (r, r)
        p = np.pad(out, pad, mode="edge")
        n = out.shape[axis]
        out = sum(k[i] * np.take(p, np.arange(i, i + n), axis=axis) for i in range(2 * r + 1))
    return out


def _erode(mask: np.ndarray, px: int) -> np.ndarray:
    img = Image.fromarray((mask * 255).astype("uint8"), "L").filter(ImageFilter.MinFilter(2 * px + 1))
    return np.asarray(img) > 127


def cutout(img: Image.Image) -> tuple[Image.Image, str]:
    """Devuelve la tarjeta recortada con alfa limpio (a escala de salida x SS) y su orientacion."""
    img = img.convert("RGBA")
    img.thumbnail((WORK_MAX, WORK_MAX), Image.Resampling.LANCZOS)  # la salida es de 560 px: no hace falta mas
    rgba = np.asarray(img)
    fg = ~background_mask(rgba)
    ys, xs = np.where(fg)
    x0, x1, y0, y1 = xs.min(), xs.max() + 1, ys.min(), ys.max() + 1
    crop = Image.fromarray(rgba[y0:y1, x0:x1])
    mask = Image.fromarray((fg[y0:y1, x0:x1] * 255).astype("uint8"), "L")

    cw, ch = CANVAS[0] - 2 * PAD, CANVAS[1] - 2 * PAD
    scale = min(cw / crop.width, ch / crop.height)
    size = (max(1, round(crop.width * scale * SS)), max(1, round(crop.height * scale * SS)))
    crop = crop.resize(size, Image.Resampling.LANCZOS)
    m = _round_corners(np.asarray(mask.resize(size, Image.Resampling.BILINEAR)) > 127)

    solid = _erode(m, ERODE)  # quita el aro claro del antialiasing original
    core = _erode(solid, ERODE)  # interior fiable para pintar la franja del borde
    rgb = np.asarray(crop.convert("RGB")).astype("float32")
    w = core.astype("float32")
    den = _blur(w, 4)
    fill = np.stack([_blur(rgb[..., c] * w, 4) for c in range(3)], axis=-1) / np.maximum(den, 1e-3)[..., None]
    use = (~core) & (den > 0.02)
    rgb[use] = fill[use]
    alpha = np.clip(_blur(solid.astype("float32") * 255, 1.2), 0, 255)
    out = np.dstack([rgb.clip(0, 255), alpha]).astype("uint8")
    orientation = "portrait" if size[1] > size[0] else "landscape"
    return Image.fromarray(out, "RGBA"), orientation


def process(path: Path, out: Path) -> str:
    card, orientation = cutout(Image.open(path))
    card = card.resize((card.width // SS, card.height // SS), Image.Resampling.LANCZOS)
    canvas = Image.new("RGBA", CANVAS, (0, 0, 0, 0))
    canvas.alpha_composite(card, ((CANVAS[0] - card.width) // 2, (CANVAS[1] - card.height) // 2))
    canvas.save(out, "WEBP", quality=90, method=6, alpha_quality=100)
    return orientation


def main() -> None:
    if not SRC_DIR.exists():
        sys.exit(f"No existe {SRC_DIR}")
    wb = load_workbook(EXCEL, read_only=True, data_only=True)
    rows = wb["Tarjetas"].iter_rows(values_only=True)
    headers = [normalize_header(h) for h in next(rows)]
    by_name = {}
    for r in rows:
        d = dict(zip(headers, r, strict=False))
        if d.get("id_tarjeta"):
            by_name[_key(str(d["nombre_de_la_tarjeta"]))] = str(d["id_tarjeta"]).strip().zfill(3)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for old in OUT_DIR.glob("*.webp"):
        old.unlink()
    manifest: dict[str, dict[str, str]] = {}
    missing: list[str] = []
    for f in sorted(SRC_DIR.iterdir()):
        cid = by_name.get(_key(f.stem))
        if cid is None:
            missing.append(f.name)
            continue
        orientation = process(f, OUT_DIR / f"{cid}.webp")
        manifest[cid] = {"file": f"/cards/{cid}.webp", "orientation": orientation}
    MANIFEST.write_text(json.dumps(manifest, indent=1, sort_keys=True), encoding="utf-8")
    size = sum(p.stat().st_size for p in OUT_DIR.glob("*.webp")) / 1024
    print(f"OK {len(manifest)} imagenes ({size:,.0f} KB). Sin tarjeta en Excel: {missing or 'ninguna'}")
    no_img = sorted(set(by_name.values()) - set(manifest))
    if no_img:
        print(f"Tarjetas sin imagen (usaran la ilustracion generada): {no_img}")


if __name__ == "__main__":
    main()
